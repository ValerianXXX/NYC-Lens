from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import duckdb
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..analyze import get_site
from ..config import metrics_config
from ..paths import project_home
from ..util import json_dumps, new_id, sha256_file, utc_now


TEMPLATE_DIR = Path(__file__).with_name("templates")


def _output_dir() -> Path:
    path = project_home() / "reports" / "samples"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _safe_site_name(site_id: str) -> str:
    return "".join(character if character.isalnum() or character in "-_" else "_" for character in site_id)


def _flatten_metrics(result: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    definitions = metrics_config().get("metrics", {})
    for catchment_name, catchment in result.get("catchments", {}).items():
        for metric_id, metric in catchment.get("metrics", {}).items():
            rows.append(
                {
                    "catchment": catchment_name,
                    "metric_id": metric_id,
                    "label_zh": definitions.get(metric_id, {}).get("label_zh", metric_id),
                    "value": metric.get("value"),
                    "text_value": metric.get("text_value"),
                    "unit": metric.get("unit"),
                    "percentile": metric.get("percentile"),
                    "source_id": metric.get("source_id"),
                    "snapshot_id": metric.get("snapshot_id"),
                    "source_year": metric.get("source_year"),
                    "measure_type": metric.get("measure_type"),
                    "confidence": metric.get("confidence"),
                    "margin_of_error": metric.get("margin_of_error"),
                    "geography_type": metric.get("geography_type"),
                    "geography_id": metric.get("geography_id"),
                    "warnings": " | ".join(metric.get("warnings") or []),
                }
            )
    return rows


def _register_report(
    conn: duckdb.DuckDBPyConnection,
    result: dict[str, Any],
    path: Path,
    format_name: str,
) -> None:
    # Report filenames are stable per site. Re-exporting must replace the
    # current registry row so an older hash never points at the overwritten
    # artifact.
    conn.execute(
        "DELETE FROM site_report WHERE site_id=? AND format=?",
        [result["site_id"], format_name],
    )
    conn.execute(
        "INSERT INTO site_report VALUES (?, ?, ?, ?, ?, ?, ?)",
        [
            new_id("report"),
            result["site_id"],
            format_name,
            str(path),
            sha256_file(path),
            utc_now(),
            result["calculation_version_id"],
        ],
    )


def export_json(conn: duckdb.DuckDBPyConnection, site_id: str) -> Path:
    result = get_site(conn, site_id)
    if result is None:
        raise KeyError(site_id)
    path = _output_dir() / f"{_safe_site_name(site_id)}.json"
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    _register_report(conn, result, path, "json")
    return path


def export_csv(conn: duckdb.DuckDBPyConnection, site_id: str) -> Path:
    result = get_site(conn, site_id)
    if result is None:
        raise KeyError(site_id)
    rows = _flatten_metrics(result)
    path = _output_dir() / f"{_safe_site_name(site_id)}.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]) if rows else ["catchment", "metric_id"])
        writer.writeheader()
        writer.writerows(rows)
    _register_report(conn, result, path, "csv")
    return path


def export_html(conn: duckdb.DuckDBPyConnection, site_id: str) -> Path:
    result = get_site(conn, site_id)
    if result is None:
        raise KeyError(site_id)
    environment = Environment(
        loader=FileSystemLoader(TEMPLATE_DIR),
        autoescape=select_autoescape(["html", "xml"]),
    )
    template = environment.get_template("site_report.html.j2")
    rows = _flatten_metrics(result)
    html = template.render(result=result, metric_rows=rows, result_json=json_dumps(result))
    path = _output_dir() / f"{_safe_site_name(site_id)}.html"
    path.write_text(html, encoding="utf-8")
    _register_report(conn, result, path, "html")
    return path


def export_xlsx(conn: duckdb.DuckDBPyConnection, site_id: str) -> Path:
    """Portable application exporter. Final delivered workbooks receive an additional artifact-tool QA pass."""
    result = get_site(conn, site_id)
    if result is None:
        raise KeyError(site_id)
    rows = _flatten_metrics(result)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "概览"
    metrics_sheet = workbook.create_sheet("指标明细")
    sources_sheet = workbook.create_sheet("来源与说明")
    navy = "16324F"
    teal = "18A999"
    light = "EAF4F2"
    amber = "F4B942"
    thin = Side(style="thin", color="D7E0E7")
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F1")
    summary["A1"] = "NYC Lens｜商圈分析概览"
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 30
    summary_rows = [
        ("地址", result.get("standardized_address")),
        ("经纬度", f"{result['coordinates']['latitude']:.6f}, {result['coordinates']['longitude']:.6f}"),
        ("行政区", result.get("borough")),
        ("BBL", (result.get("parcel") or {}).get("bbl")),
        ("Census Tract", result.get("census_tract")),
        ("H3 Res 10", result.get("h3_resolution_10")),
        ("商业类型", result.get("business_type")),
        ("计算版本", result.get("calculation_version_id")),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=3):
        summary.cell(row_index, 1, label).font = Font(bold=True, color=navy)
        summary.cell(row_index, 2, value)
    summary["A13"] = "重要说明"
    summary["A13"].font = Font(bold=True, color=navy)
    summary.merge_cells("A14:F16")
    summary["A14"] = (
        "ACS为区域调查估计并保留误差范围；LODES统计岗位而非独立自然人；游客、夜生活与客群判断为开放数据代理。"
        "所有健康、收入、语言与族裔信息仅描述公开地理范围，不代表具体个人、家庭或建筑。"
    )
    summary["A14"].alignment = Alignment(wrap_text=True, vertical="top")
    summary["A14"].fill = PatternFill("solid", fgColor="FFF4D6")
    summary.column_dimensions["A"].width = 22
    for column in range(2, 7):
        summary.column_dimensions[get_column_letter(column)].width = 18

    headers = list(rows[0]) if rows else ["catchment", "metric_id"]
    metrics_sheet.append(headers)
    for row in rows:
        metrics_sheet.append([row.get(header) for header in headers])
    metrics_sheet.freeze_panes = "A2"
    metrics_sheet.auto_filter.ref = metrics_sheet.dimensions
    metrics_sheet.sheet_view.showGridLines = False
    for cell in metrics_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for row in metrics_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {2, 13, 14})
    for index, width in enumerate([16, 34, 14, 22, 14, 20, 22, 18, 28, 12, 16, 22, 20, 45], start=1):
        metrics_sheet.column_dimensions[get_column_letter(index)].width = width
    if "value" in headers:
        value_col = headers.index("value") + 1
        for cell in metrics_sheet.iter_cols(min_col=value_col, max_col=value_col, min_row=2):
            for item in cell:
                item.number_format = "#,##0.00"
    if "confidence" in headers:
        confidence_col = headers.index("confidence") + 1
        for cell in metrics_sheet.iter_cols(min_col=confidence_col, max_col=confidence_col, min_row=2):
            for item in cell:
                item.number_format = "0%"

    sources_sheet.sheet_view.showGridLines = False
    sources_sheet.append(["source_id", "snapshot_id", "source_year", "measure_type", "说明"])
    unique_sources: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        key = (row["source_id"], row["snapshot_id"], row["source_year"], row["measure_type"])
        unique_sources[key] = row
    for key, row in unique_sources.items():
        sources_sheet.append([*key, row.get("warnings")])
    for cell in sources_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.font = Font(color="FFFFFF", bold=True)
    sources_sheet.freeze_panes = "A2"
    for column, width in zip("ABCDE", [26, 40, 24, 30, 70]):
        sources_sheet.column_dimensions[column].width = width
    for row in sources_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)

    path = _output_dir() / f"{_safe_site_name(site_id)}.xlsx"
    workbook.save(path)
    _register_report(conn, result, path, "xlsx")
    return path


def export_all(
    conn: duckdb.DuckDBPyConnection,
    site_id: str,
    *,
    include_xlsx: bool = True,
) -> dict[str, str]:
    exports = {
        "json": str(export_json(conn, site_id)),
        "csv": str(export_csv(conn, site_id)),
        "html": str(export_html(conn, site_id)),
    }
    if include_xlsx:
        exports["xlsx"] = str(export_xlsx(conn, site_id))
    return exports
