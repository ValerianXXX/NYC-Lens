from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..comparison import get_comparison
from ..paths import project_home
from ..util import new_id, sha256_file, utc_now


def _output_dir() -> Path:
    path = project_home() / "reports" / "comparisons"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _safe_name(value: str) -> str:
    return "".join(character if character.isalnum() or character in "-_" else "_" for character in value)


def _rows(result: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for metric_id, comparison in result.get("metric_comparisons", {}).items():
        for item in comparison.get("values", []):
            rows.append(
                {
                    "metric_id": metric_id,
                    "label_zh": comparison.get("label_zh"),
                    "site_id": item.get("site_id"),
                    "address": item.get("address"),
                    "value": item.get("value"),
                    "text_value": item.get("text_value"),
                    "unit": item.get("unit"),
                    "high_to_low_rank": item.get("high_to_low_rank"),
                    "within_batch_high_value_percentile": item.get(
                        "within_batch_high_value_percentile"
                    ),
                    "source_id": item.get("source_id"),
                    "snapshot_id": item.get("snapshot_id"),
                    "source_year": item.get("source_year"),
                    "measure_type": item.get("measure_type"),
                    "confidence": item.get("confidence"),
                    "margin_of_error": item.get("margin_of_error"),
                    "warnings": " | ".join(item.get("warnings") or []),
                }
            )
    return rows


def _register(
    conn: duckdb.DuckDBPyConnection,
    result: dict[str, Any],
    path: Path,
    format_name: str,
) -> None:
    conn.execute(
        "DELETE FROM comparison_report WHERE comparison_id=? AND format=?",
        [result["comparison_id"], format_name],
    )
    conn.execute(
        "INSERT INTO comparison_report VALUES (?, ?, ?, ?, ?, ?, ?)",
        [
            new_id("comparison_report"),
            result["comparison_id"],
            format_name,
            str(path),
            sha256_file(path),
            utc_now(),
            result["calculation_version_id"],
        ],
    )


def export_comparison_json(conn: duckdb.DuckDBPyConnection, comparison_id: str) -> Path:
    result = get_comparison(conn, comparison_id)
    if result is None:
        raise KeyError(comparison_id)
    path = _output_dir() / f"{_safe_name(comparison_id)}.json"
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    _register(conn, result, path, "json")
    return path


def export_comparison_csv(conn: duckdb.DuckDBPyConnection, comparison_id: str) -> Path:
    result = get_comparison(conn, comparison_id)
    if result is None:
        raise KeyError(comparison_id)
    rows = _rows(result)
    path = _output_dir() / f"{_safe_name(comparison_id)}.csv"
    headers = list(rows[0]) if rows else ["metric_id", "site_id"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    _register(conn, result, path, "csv")
    return path


def export_comparison_xlsx(conn: duckdb.DuckDBPyConnection, comparison_id: str) -> Path:
    result = get_comparison(conn, comparison_id)
    if result is None:
        raise KeyError(comparison_id)
    rows = _rows(result)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "比较概览"
    detail = workbook.create_sheet("同口径指标")
    sources = workbook.create_sheet("来源与限制")
    navy = "16324F"
    teal = "18A999"
    thin = Side(style="thin", color="D7E0E7")
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:F1")
    overview["A1"] = "NYC Lens｜多地址同口径比较"
    overview["A1"].fill = PatternFill("solid", fgColor=navy)
    overview["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    overview["A1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 30
    overview.append([])
    overview.append(["比较 ID", result["comparison_id"]])
    overview.append(["商业类型", result["business_type"]])
    overview.append(["比较商圈", result["comparison_catchment"]])
    overview.append(["计算版本", result["calculation_version_id"]])
    overview.append(["地址数量", len(result["sites"])])
    overview.append([])
    overview.append(["站点", "标准地址", "州", "County GEOID", "BBL", "缺失项数量"])
    for site in result["sites"]:
        overview.append(
            [
                site["site_id"],
                site["standardized_address"],
                site.get("state_abbr"),
                site.get("county_geoid"),
                site.get("parcel_bbl"),
                len(site.get("missing_data") or []),
            ]
        )
    warning_row = overview.max_row + 2
    overview.cell(warning_row, 1, "重要说明").font = Font(bold=True, color=navy)
    overview.merge_cells(start_row=warning_row + 1, start_column=1, end_row=warning_row + 4, end_column=6)
    overview.cell(warning_row + 1, 1, "\n".join(result.get("warnings") or []))
    overview.cell(warning_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    overview.cell(warning_row + 1, 1).fill = PatternFill("solid", fgColor="FFF4D6")
    for column, width in enumerate([28, 55, 10, 16, 18, 14], start=1):
        overview.column_dimensions[get_column_letter(column)].width = width

    headers = list(rows[0]) if rows else ["metric_id", "site_id"]
    detail.append(headers)
    for row in rows:
        detail.append([row.get(header) for header in headers])
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    detail.sheet_view.showGridLines = False
    for cell in detail[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {2, 4, 16})
    for index, width in enumerate(
        [34, 28, 30, 55, 16, 22, 18, 15, 18, 28, 38, 24, 30, 14, 16, 65],
        start=1,
    ):
        detail.column_dimensions[get_column_letter(index)].width = width

    sources.append(["source_id", "snapshot_id", "source_year", "measure_type", "说明"])
    unique: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        key = (row["source_id"], row["snapshot_id"], row["source_year"], row["measure_type"])
        unique[key] = row
    for key, row in unique.items():
        sources.append([*key, row.get("warnings")])
    sources.freeze_panes = "A2"
    sources.sheet_view.showGridLines = False
    for cell in sources[1]:
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.font = Font(color="FFFFFF", bold=True)
    for column, width in zip("ABCDE", [28, 40, 28, 35, 75]):
        sources.column_dimensions[column].width = width
    for row in sources.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    path = _output_dir() / f"{_safe_name(comparison_id)}.xlsx"
    workbook.save(path)
    _register(conn, result, path, "xlsx")
    return path


def export_comparison_all(
    conn: duckdb.DuckDBPyConnection,
    comparison_id: str,
) -> dict[str, str]:
    return {
        "json": str(export_comparison_json(conn, comparison_id)),
        "csv": str(export_comparison_csv(conn, comparison_id)),
        "xlsx": str(export_comparison_xlsx(conn, comparison_id)),
    }
